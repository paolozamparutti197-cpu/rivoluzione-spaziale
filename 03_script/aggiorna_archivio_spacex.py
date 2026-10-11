"""Ingresso Python della pipeline archivio: workbook sempre in sola lettura.

Equivalente allo script PowerShell esistente, senza dipendere da execution policy.
Non salva o ricalcola Excel e conserva le impronte delle fonti.
"""
from pathlib import Path
from datetime import date, datetime, time
from zoneinfo import ZoneInfo
import hashlib,json,os,shutil,subprocess,sys,argparse
import openpyxl
from openpyxl.utils.datetime import to_excel,CALENDAR_MAC_1904

ROOT=Path(__file__).resolve().parents[1]

def sheet(relative,name,header_row,date_columns):
    book=openpyxl.load_workbook(ROOT/relative,read_only=True,data_only=True)
    try:
        if book.epoch==CALENDAR_MAC_1904:raise ValueError('Date Excel 1904 non supportate')
        if name not in book.sheetnames:raise ValueError(f'Foglio mancante: {name}')
        rows=book[name].iter_rows(min_row=header_row)
        headers=[str(c.value) if c.value is not None else None for c in next(rows)]
        result=[]
        for rownum,row in enumerate(rows,header_row+1):
            record={'source_row':rownum}
            for key,cell in zip(headers,row):
                if not key:continue
                value=cell.value
                if cell.data_type=='e':raise ValueError(f'Errore Excel {relative}/{name}/{cell.coordinate}: {value}')
                if isinstance(value,(datetime,date,time)):
                    if key in date_columns and not isinstance(value,time):value=value.strftime('%Y-%m-%d')
                    elif key=='ora_utc':value=value.strftime('%H:%M:%S')
                    else:value=to_excel(value,book.epoch)
                elif value is not None and key in date_columns and isinstance(value,(int,float)):
                    from openpyxl.utils.datetime import from_excel
                    value=from_excel(value,book.epoch).strftime('%Y-%m-%d')
                elif value is not None and key=='ora_utc' and isinstance(value,(int,float)):
                    from openpyxl.utils.datetime import from_excel
                    value=from_excel(value,book.epoch).strftime('%H:%M:%S')
                if isinstance(value,bool):value=int(value)
                record[key]=value
            if record.get('id_lancio') or record.get('Flight') or record.get('id_sito'):result.append(record)
        return result
    finally:book.close()

def update_archive(imported=None,node=None):
    sources=['01_workbook/lanci_spacex.xlsx','01_workbook/sviluppo_starship.xlsx','monografie/starship/voli.json','sezioni/storico-lanci.html']
    def fingerprints():return {p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in sources}
    hashes=fingerprints()
    payload={'imported':imported or datetime.now(ZoneInfo('Europe/Rome')).date().isoformat(),'hashes':hashes,
        'launches':sheet(sources[0],'elenco',1,['data','data_evento','anno_mese']),
        'recoveries':sheet(sources[0],'recuperi_veicoli',1,['data_recupero']),
        'sites':sheet(sources[0],'siti',1,[]),
        'tests':sheet(sources[1],'Voli integrati',4,[]),
        'monograph':json.loads((ROOT/sources[2]).read_text(encoding='utf8')),
        'legacy_html':(ROOT/sources[3]).read_text(encoding='utf8')}
    if hashes!=fingerprints():raise RuntimeError('Fonte cambiata durante la lettura: importazione interrotta')
    scratch=ROOT/'99_temporanei/archivio110';scratch.mkdir(parents=True,exist_ok=True)
    raw=scratch/'fonti.json';raw.write_text(json.dumps(payload,ensure_ascii=False),encoding='utf8')
    node=node or shutil.which('node') or str(Path(os.environ['USERPROFILE'])/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe')
    env=dict(os.environ,SPACEX_PYTHON=sys.executable)
    subprocess.check_call([node,str(ROOT/'03_script/costruisci_archivio_spacex.mjs'),str(raw)],cwd=ROOT,env=env)
    for p in sources[:2]:
        if hashlib.sha256((ROOT/p).read_bytes()).hexdigest()!=hashes[p]:raise RuntimeError(f'Workbook cambiato durante la costruzione: {p}')

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--data-importazione');parser.add_argument('--node')
    args=parser.parse_args();update_archive(args.data_importazione,args.node)
